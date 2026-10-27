#!/usr/bin/env python3
"""Build the 10-slide Healthcare Appointment System presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

BASE = "/home/hatch/workspace/job-search/project-15-healthcare-appointment"


def slide(title, bullets, image=None):
    s = prs.slides.add_slide(BLANK)
    tx = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    if image:
        s.shapes.add_picture(image, Inches(0.6), Inches(1.4), width=Inches(6.2))
        bx, bw = Inches(7.1), Inches(5.6)
    else:
        bx, bw = Inches(0.6), Inches(12.1)
    box = s.shapes.add_textbox(bx, Inches(1.4), bw, Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(22)
        p.space_after = Pt(14)
        p.level = 0
    return s


slide("Healthcare Patient and Appointment Management System", [
    "Spring Boot backend for patients, providers, and appointments",
    "By Ramesh Poluri — Java developer",
    "GitHub: github.com/rpoluri9/healthcare-patient-appointment",
    "All demo patients and providers are fictional",
])

slide("The idea", [
    "A clinic backend that books appointments without double-booking anyone",
    "Patients, providers, and appointments as clean REST resources",
    "Scheduling rules enforced on the server, not trusted to the client",
    "Roles and ownership so patients only ever see their own records",
])

slide("Architecture", [
    "Controllers -> Services -> Repositories -> H2 database",
    "Spring Security filter chain in front of every endpoint",
    "Spring Data JPA with overlap queries for conflict detection",
    "PostgreSQL documented as the production database option",
], image=f"{BASE}/docs/architecture.png")

slide("Patients and providers", [
    "Patient registration with validation: names, past birth date, phone, email",
    "Provider profiles: specialization, working days, working hours, visit length",
    "List endpoints return summaries only — no sensitive fields leak",
    "Full patient records only to staff, admin, or the patient themselves",
])

slide("Scheduling rules that prevent conflicts", [
    "No bookings in the past",
    "Day must be a provider working day; slot must fit working hours",
    "End time is derived from the provider's visit length",
    "Overlapping provider booking -> HTTP 409",
    "Overlapping patient booking -> HTTP 409",
], image=f"{BASE}/docs/appointment-lifecycle.png")

slide("Appointment lifecycle", [
    "BOOKED -> CANCELLED (staff, admin, or the patient)",
    "BOOKED -> COMPLETED (staff or admin only)",
    "Cancelled slots free up immediately for rebooking",
    "Search by patient, provider, status, or date range",
])

slide("Authentication and roles", [
    "Local HTTP Basic auth, BCrypt password hashes, no plain-text passwords",
    "ADMIN: everything. STAFF: manage records. PATIENT: own records only",
    "Patient accounts link to a patient record for ownership checks",
    "No external identity provider — and none is claimed",
], image=f"{BASE}/docs/auth-role-flow.png")

slide("Secure data handling", [
    "List views never include birth dates, phone, email, or addresses",
    "Logs carry entity IDs only — never patient names or contact details",
    "Centralized errors: one clean shape, no stack traces to clients",
    "Demo data is fictional; production use needs real compliance review",
])

slide("Tests and verification", [
    "Unit tests: every scheduling rule, plus patient and provider services",
    "Integration tests: full booking flow with real HTTP Basic auth",
    "Security tests: 401 unauthenticated, 403 wrong role, 403 other patient",
    "Live boot check: register, add provider, book, conflict, cancel, search",
])

slide("Try it", [
    "Run: mvn spring-boot:run — service on http://localhost:8080",
    "Demo logins: admin/admin123, staff/staff123, patient1/patient123",
    "Postman collection included for every endpoint",
    "Thank you — questions welcome",
])

out = "/home/hatch/workspace/your_files/healthcare-appointment-presentation.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
